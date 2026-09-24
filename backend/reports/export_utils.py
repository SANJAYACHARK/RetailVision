from io import BytesIO
from datetime import datetime

from django.http import HttpResponse

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


# ============================================================
# COMMON HELPERS
# ============================================================

def safe_value(value):
    if value is None:
        return ""

    return value


def timestamp_string():
    return datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )


# ============================================================
# EXCEL EXPORT
# ============================================================

def generate_excel_response(
    *,
    title,
    headers,
    rows,
    filename,
):

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = title[:31]


    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=len(headers),
    )

    title_cell = worksheet.cell(
        row=1,
        column=1,
        value=f"RetailVision - {title}",
    )

    title_cell.font = Font(
        bold=True,
        size=16,
    )

    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    worksheet.row_dimensions[1].height = 28


    # --------------------------------------------------------
    # EXPORT DATE
    # --------------------------------------------------------

    worksheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=len(headers),
    )

    date_cell = worksheet.cell(
        row=2,
        column=1,
        value=(
            "Generated: "
            + datetime.now().strftime(
                "%d %b %Y, %I:%M %p"
            )
        ),
    )

    date_cell.alignment = Alignment(
        horizontal="center",
    )

    date_cell.font = Font(
        italic=True,
        size=10,
    )


    # --------------------------------------------------------
    # HEADER
    # --------------------------------------------------------

    header_row = 4

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="E8D65A",
    )

    for column_index, header in enumerate(
        headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=header_row,
            column=column_index,
            value=header,
        )

        cell.font = Font(
            bold=True,
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


    # --------------------------------------------------------
    # DATA
    # --------------------------------------------------------

    for row_index, row in enumerate(
        rows,
        start=header_row + 1,
    ):

        for column_index, value in enumerate(
            row,
            start=1,
        ):

            cell = worksheet.cell(
                row=row_index,
                column=column_index,
                value=safe_value(value),
            )

            cell.alignment = Alignment(
                vertical="top",
            )


    # --------------------------------------------------------
    # COLUMN WIDTH
    # --------------------------------------------------------

    for column_index in range(
        1,
        len(headers) + 1,
    ):

        max_length = 0

        column_letter = get_column_letter(
            column_index
        )

        for cell in worksheet[
            column_letter
        ]:

            value = cell.value

            if value is None:
                continue

            length = len(
                str(value)
            )

            if length > max_length:
                max_length = length

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 3,
            35,
        )


    # --------------------------------------------------------
    # FREEZE HEADER
    # --------------------------------------------------------

    worksheet.freeze_panes = "A5"


    # --------------------------------------------------------
    # RESPONSE
    # --------------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )

    return response


# ============================================================
# PDF EXPORT
# ============================================================

def generate_pdf_response(
    *,
    title,
    headers,
    rows,
    filename,
):

    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )

    styles = getSampleStyleSheet()

    elements = []


    # --------------------------------------------------------
    # TITLE
    # --------------------------------------------------------

    elements.append(
        Paragraph(
            f"<b>RetailVision - {title}</b>",
            styles["Title"],
        )
    )

    elements.append(
        Spacer(
            1,
            5 * mm,
        )
    )

    elements.append(
        Paragraph(
            (
                "Generated: "
                + datetime.now().strftime(
                    "%d %b %Y, %I:%M %p"
                )
            ),
            styles["Normal"],
        )
    )

    elements.append(
        Spacer(
            1,
            5 * mm,
        )
    )


    # --------------------------------------------------------
    # TABLE DATA
    # --------------------------------------------------------

    table_data = [
        headers
    ]

    for row in rows:
        table_data.append(
            [
                safe_value(value)
                for value in row
            ]
        )


    table = Table(
        table_data,
        repeatRows=1,
    )


    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor(
                        "#E8D65A"
                    ),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor(
                        "#252523"
                    ),
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "FONTNAME",
                    (0, 1),
                    (-1, -1),
                    "Helvetica",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, 0),
                    "CENTER",
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.HexColor(
                        "#D6D3C8"
                    ),
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor(
                            "#FAF9F4"
                        ),
                    ],
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
            ]
        )
    )

    elements.append(
        table
    )

    document.build(
        elements
    )

    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/pdf",
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )

    return response